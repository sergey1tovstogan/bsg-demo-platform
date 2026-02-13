"""Extract text from Temenos Active and Future Components.docx into a text file."""
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)

DOCX_PATH = Path(r"c:\Users\apostolos.georgas\OneDrive - Temenos\Desktop\Temenos\Technical Stuff\Docs\Temenos Active and Future Components.docx")
OUT_PATH = Path(__file__).resolve().parent.parent / "temenos_components_extract.txt"

def main():
    if not DOCX_PATH.exists():
        print(f"File not found: {DOCX_PATH}", file=sys.stderr)
        sys.exit(1)
    doc = Document(DOCX_PATH)
    lines = []
    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            lines.append(" | ".join(cells))
    out_text = "\n".join(lines)
    OUT_PATH.write_text(out_text, encoding="utf-8")
    print(f"Written {len(lines)} lines to {OUT_PATH}")

if __name__ == "__main__":
    main()
