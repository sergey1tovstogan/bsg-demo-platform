#!/usr/bin/env python3
"""
Script to extract paragraphs from Word document and import into PostgreSQL database.
"""

import sys
import os
from docx import Document
import psycopg2
from psycopg2.extras import execute_values

# Database connection parameters
DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'database': 'bsg_demo',
    'user': 'postgres',
    'password': 'postgres'
}

def extract_paragraphs_from_docx(file_path):
    """Extract paragraphs from Word document."""
    print(f"Reading document: {file_path}")
    doc = Document(file_path)
    
    paragraphs = []
    paragraph_number = 1
    
    for para in doc.paragraphs:
        text = para.text.strip()
        # Skip empty paragraphs
        if text:
            paragraphs.append((paragraph_number, text))
            paragraph_number += 1
    
    print(f"Extracted {len(paragraphs)} paragraphs")
    return paragraphs

def import_paragraphs_to_db(paragraphs):
    """Import paragraphs into PostgreSQL database."""
    try:
        conn = psycopg2.connect(**DB_CONFIG)
        cur = conn.cursor()
        
        # Clear existing data
        cur.execute('DELETE FROM "component-security".paragraphs;')
        print("Cleared existing paragraphs")
        
        # Insert new paragraphs
        insert_query = '''
            INSERT INTO "component-security".paragraphs (paragraph_number, paragraph_content)
            VALUES %s
            ON CONFLICT (paragraph_number) DO UPDATE 
            SET paragraph_content = EXCLUDED.paragraph_content;
        '''
        
        execute_values(cur, insert_query, paragraphs)
        conn.commit()
        
        print(f"Successfully imported {len(paragraphs)} paragraphs into database")
        
        # Verify import
        cur.execute('SELECT COUNT(*) FROM "component-security".paragraphs;')
        count = cur.fetchone()[0]
        print(f"Total paragraphs in database: {count}")
        
        cur.close()
        conn.close()
        
    except Exception as e:
        print(f"Error importing to database: {e}")
        sys.exit(1)

def main():
    docx_path = "/Users/sergeytovstogan/Documents_BSG/Temenos Security Framework Overview_ST_SEP-ORIGINAL 2025.docx"
    
    if not os.path.exists(docx_path):
        print(f"Error: File not found: {docx_path}")
        sys.exit(1)
    
    # Extract paragraphs
    paragraphs = extract_paragraphs_from_docx(docx_path)
    
    if not paragraphs:
        print("No paragraphs found in document")
        sys.exit(1)
    
    # Import to database
    import_paragraphs_to_db(paragraphs)
    
    print("\nImport completed successfully!")

if __name__ == "__main__":
    main()

